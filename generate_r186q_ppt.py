#!/usr/bin/env python3
"""Generate a 1-slide PPT: case evidence summary for TNNI3 p.R186Q."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
BLUE = RGBColor(0x2E, 0x5E, 0xAA)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xFA)
RED = RGBColor(0xC0, 0x3A, 0x2B)
ORANGE = RGBColor(0xD9, 0x7B, 0x29)
TEAL = RGBColor(0x1F, 0x8A, 0x70)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_RED = RGBColor(0xFB, 0xEA, 0xE8)
LIGHT_ORANGE = RGBColor(0xFC, 0xF1, 0xE3)
LIGHT_TEAL = RGBColor(0xE4, 0xF4, 0xF0)
ROW_ALT = RGBColor(0xF4, 0xF7, 0xFC)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def set_font(run, size, color=DARK, bold=False):
    f = run.font
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_box(slide, x, y, w, h, fill=None, line=None, line_w=None,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=3):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for line in lines:
        para_opts = {}
        if isinstance(line, tuple) and len(line) == 2 and isinstance(line[1], dict):
            runs, para_opts = line
        else:
            runs = line
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para_opts.get("align", align)
        p.space_after = Pt(para_opts.get("space_after", space_after))
        for text, size, color, bold in runs:
            r = p.add_run()
            r.text = text
            set_font(r, size, color, bold)
    return tb


s = prs.slides.add_slide(blank)

# ---------- header ----------
add_box(s, 0, 0, SW, Inches(0.92), fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
add_box(s, 0, Inches(0.92), SW, Inches(0.05), fill=RED, shape=MSO_SHAPE.RECTANGLE)
add_text(s, Inches(0.5), Inches(0.1), Inches(11.0), Inches(0.5),
         [[("TNNI3 p.Arg186Gln（R186Q）病例证据汇总", 24, WHITE, True)]])
add_text(s, Inches(0.5), Inches(0.56), Inches(11.0), Inches(0.3),
         [[("Case-Level Evidence for TNNI3 R186Q (c.557G>A) in Cardiomyopathy", 11, RGBColor(0xB8, 0xC6, 0xDD), False)]])

# ---------- variant info strip ----------
info_y = Inches(1.08)
add_box(s, Inches(0.35), info_y, Inches(12.63), Inches(0.6), fill=LIGHT_BLUE)
add_text(s, Inches(0.6), info_y + Inches(0.06), Inches(12.2), Inches(0.5), [
    [("变异信息：", 10.5, BLUE, True),
     ("c.557G>A（8 号外显子）· rs397516357 · 位于 C 端移动区（164–210）　｜　", 10.5, DARK, False),
     ("分类：", 10.5, BLUE, True),
     ("ClinVar 致病/可能致病（11 家实验室）；OMGL Class 5 确定致病　｜　", 10.5, DARK, False),
     ("人群频率：", 10.5, BLUE, True),
     ("ExAC（6 万人）未检出；gnomAD v3.1.1 仅 1 名携带者", 10.5, DARK, False)],
], anchor=MSO_ANCHOR.MIDDLE)

# ---------- case evidence table ----------
add_text(s, Inches(0.35), Inches(1.82), Inches(8.0), Inches(0.32),
         [[("已报道病例 / 家系证据", 15, NAVY, True)]])

tbl_x, tbl_y = Inches(0.35), Inches(2.18)
col_w = [Inches(2.5), Inches(2.35), Inches(7.78)]
headers = ["来源 / 年份", "病例 / 家系", "关键临床发现"]
rows_data = [
    ("Mogensen 等, JACC 2004\n英国 748 家系 HCM 队列",
     "家系 H136（该队列 23 个 TNNI3 阳性家系之一）",
     "先证者 29 岁诊断，重度非对称室间隔肥厚达 35 mm；随访 13 年出现左室壁变薄、心腔扩张（HCM → 终末期/扩张样演变）"),
    ("Moon 等, Heart 2005\n同队列 CMR 随访研究",
     "家系 H136，≥4 名基因阳性成员",
     "2 人 40 余岁死于进展性心衰；1 例 6 年内出现 LVH + 晕厥（LGE 11%）；1 例 LVH>30 mm、运动试验异常，后室壁变薄扩张（LGE 高达 48%，广泛纤维化）"),
    ("Wang 等, J Med Genet 2015\n中国 GeneID 队列",
     "1 个 HCM + 房颤大家系",
     "全外显子测序发现 R186Q 与疾病在家系内共分离，>1583 名对照未检出 —— 提示 R186Q 可同时导致 HCM 与房颤（AF）"),
    ("临床心血管病杂志 2018\n中国汉族家族性 HCM",
     "1 家系 3 例 HCM 患者",
     "先证者无症状、仅体检心电图异常；其父胸闷、研究随访期间猝死；祖母胸闷气促明显 —— 同一家系内表型异质性显著"),
    ("俄罗斯病例报告 2024",
     "青少年早期 HCM 1 例",
     "杂合 R186Q；超声应变成像示左室心肌节段形变减低，为肥厚前早期改变提供影像学证据"),
    ("检测队列汇总\n（Atlas of Cardiac Genetic Variation）",
     "OMGL：2/3135 HCM\nLMM：3/2912 HCM",
     "两大临床检测实验室在 HCM 队列中独立多次检出并均判为致病；DCM/对照队列未见富集"),
]

header_h = Inches(0.38)
row_hs = [Inches(0.62), Inches(0.74), Inches(0.56), Inches(0.56), Inches(0.44), Inches(0.5)]

# header row
cx = tbl_x
for i, htxt in enumerate(headers):
    add_box(s, cx, tbl_y, col_w[i], header_h, fill=BLUE, shape=MSO_SHAPE.RECTANGLE, line=WHITE, line_w=Pt(1))
    add_text(s, cx + Inches(0.12), tbl_y, col_w[i] - Inches(0.24), header_h,
             [[(htxt, 11, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    cx += col_w[i]

cy = tbl_y + header_h
for r, (src, fam, finding) in enumerate(rows_data):
    bg = ROW_ALT if r % 2 == 0 else WHITE
    rh = row_hs[r]
    cx = tbl_x
    for c, txt in enumerate((src, fam, finding)):
        add_box(s, cx, cy, col_w[c], rh, fill=bg, shape=MSO_SHAPE.RECTANGLE,
                line=RGBColor(0xD5, 0xDE, 0xEA), line_w=Pt(0.75))
        color = NAVY if c == 0 else DARK
        bold = c == 0
        lines = [[(seg, 9.5, color, bold)] for seg in txt.split("\n")]
        add_text(s, cx + Inches(0.12), cy + Inches(0.03), col_w[c] - Inches(0.24), rh - Inches(0.06),
                 lines, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
        cx += col_w[c]
    cy += rh

# ---------- bottom boxes ----------
by = cy + Inches(0.12)
bh = Inches(7.5) - by - Inches(0.12)

add_box(s, Inches(0.35), by, Inches(6.15), bh, fill=LIGHT_TEAL)
add_text(s, Inches(0.6), by + Inches(0.07), Inches(5.7), Inches(0.3),
         [[("功能学 / 机制证据（J Cardiovasc Aging 2022）", 11.5, TEAL, True)]])
add_text(s, Inches(0.6), by + Inches(0.38), Inches(5.7), bh - Inches(0.45), [
    [("▪ ", 9.5, TEAL, True), ("R186Q 敲入小鼠再现心肌肥厚表型（体内致病性证据）", 9.5, DARK, False)],
    [("▪ ", 9.5, TEAL, True), ("机制：cTnI–EGFR 结合减弱 → FASN 上调 → 脂肪酸代谢异常；FASN 抑制剂 C75 在模型中可改善表型", 9.5, DARK, False)],
], space_after=3)

add_box(s, Inches(6.7), by, Inches(6.28), bh, fill=LIGHT_RED)
add_text(s, Inches(6.95), by + Inches(0.07), Inches(5.85), Inches(0.3),
         [[("要点小结", 11.5, RED, True)]])
add_text(s, Inches(6.95), by + Inches(0.38), Inches(5.85), bh - Inches(0.45), [
    [("▪ ", 9.5, RED, True), ("表型：HCM 为主（可重度肥厚 35 mm），可进展为室壁变薄/扩张的终末期表型；伴房颤、心衰死亡与猝死风险", 9.5, DARK, False)],
    [("▪ ", 9.5, RED, True), ("多国家系共分离 + 人群中极罕见 + 动物模型验证 → 致病性证据链完整；同密码子 R186G 亦有致病报道（HCM/LVNC 混合表型）", 9.5, DARK, False)],
], space_after=3)

prs.save("/workspace/TNNI3_R186Q病例证据.pptx")
print("saved")
