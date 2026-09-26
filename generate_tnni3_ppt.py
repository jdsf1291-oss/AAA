#!/usr/bin/env python3
"""Generate a 3-slide PPT: TNNI3 gene mutations and cardiomyopathy."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette ----------
NAVY = RGBColor(0x1B, 0x2A, 0x4A)      # deep navy background accents
BLUE = RGBColor(0x2E, 0x5E, 0xAA)      # primary blue
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xFA)
RED = RGBColor(0xC0, 0x3A, 0x2B)       # HCM
ORANGE = RGBColor(0xD9, 0x7B, 0x29)    # RCM
TEAL = RGBColor(0x1F, 0x8A, 0x70)      # DCM
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_RED = RGBColor(0xFB, 0xEA, 0xE8)
LIGHT_ORANGE = RGBColor(0xFC, 0xF1, 0xE3)
LIGHT_TEAL = RGBColor(0xE4, 0xF4, 0xF0)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def set_font(run, size, color=DARK, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.italic = italic
    f.name = FONT
    # ensure East-Asian font is set too
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_box(slide, x, y, w, h, fill=None, line=None, line_w=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    """lines: list of lists of (text, size, color, bold) run tuples, or dict for para opts."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for line in lines:
        para_opts = {}
        if isinstance(line, tuple) and len(line) == 2 and isinstance(line[1], dict):
            runs, para_opts = line
        else:
            runs = line
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para_opts.get("align", align)
        p.space_after = Pt(para_opts.get("space_after", space_after))
        p.space_before = Pt(para_opts.get("space_before", 0))
        if "line_spacing" in para_opts:
            p.line_spacing = para_opts["line_spacing"]
        for text, size, color, bold in runs:
            r = p.add_run()
            r.text = text
            set_font(r, size, color, bold)
    return tb


def header(slide, title_cn, subtitle, page_no):
    add_box(slide, 0, 0, SW, Inches(1.05), fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    add_box(slide, 0, Inches(1.05), SW, Inches(0.06), fill=RED, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, Inches(0.55), Inches(0.14), Inches(10.5), Inches(0.55),
             [[(title_cn, 26, WHITE, True)]])
    add_text(slide, Inches(0.55), Inches(0.63), Inches(10.5), Inches(0.35),
             [[(subtitle, 12.5, RGBColor(0xB8, 0xC6, 0xDD), False)]])
    add_text(slide, Inches(12.35), Inches(0.3), Inches(0.7), Inches(0.5),
             [[(f"{page_no} / 3", 13, RGBColor(0xB8, 0xC6, 0xDD), False)]], align=PP_ALIGN.RIGHT)


def bullet_runs(label, text, label_color, size=12.5):
    return [("▪ ", size, label_color, True), (label, size, label_color, True), (text, size, DARK, False)]


# ============================================================
# Slide 1 — TNNI3 gene / cTnI protein and domain structure
# ============================================================
s = prs.slides.add_slide(blank)
header(s, "TNNI3 基因与心肌肌钙蛋白 I（cTnI）", "TNNI3 Gene and Cardiac Troponin I: Structure and Function", 1)

# Left panel: gene & function
add_box(s, Inches(0.45), Inches(1.4), Inches(6.1), Inches(2.6), fill=LIGHT_BLUE)
add_text(s, Inches(0.75), Inches(1.55), Inches(5.6), Inches(0.4),
         [[("基因与蛋白概况", 16, BLUE, True)]])
add_text(s, Inches(0.75), Inches(2.0), Inches(5.6), Inches(1.95), [
    bullet_runs("定位：", "染色体 19q13.4，编码心肌肌钙蛋白 I（cTnI，210 aa）", BLUE),
    bullet_runs("功能：", "肌钙蛋白复合物（TnI–TnT–TnC）的抑制亚基", BLUE),
    bullet_runs("低钙状态：", "cTnI 结合肌动蛋白–原肌球蛋白，抑制肌动球蛋白 ATP 酶 → 阻止收缩（维持舒张）", BLUE),
    bullet_runs("高钙状态：", "Ca²⁺ 结合 TnC → cTnI 开关区与 TnC 结合，解除抑制 → 收缩", BLUE),
], space_after=7)

# Right panel: mutation overview
add_box(s, Inches(6.8), Inches(1.4), Inches(6.1), Inches(2.6), fill=LIGHT_RED)
add_text(s, Inches(7.1), Inches(1.55), Inches(5.6), Inches(0.4),
         [[("突变谱概览", 16, RED, True)]])
add_text(s, Inches(7.1), Inches(2.0), Inches(5.6), Inches(1.95), [
    bullet_runs("表型谱：", "同一基因可致 HCM、RCM、DCM 三种心肌病", RED),
    bullet_runs("占比：", "约 5% 的 HCM；遗传性 RCM 的最常见致病基因", RED),
    bullet_runs("热点区：", "显性致病突变高度集中于 C 端 1/3（残基 141–209）", RED),
    bullet_runs("遗传方式：", "多为常染色体显性错义突变；双等位（隐性）截短变异见于婴幼儿重症 DCM", RED),
], space_after=7)

# Domain diagram
add_text(s, Inches(0.45), Inches(4.25), Inches(8.0), Inches(0.4),
         [[("cTnI 功能结构域（1–210 aa）— 致病突变热点区", 16, NAVY, True)]])

bar_y = Inches(4.85)
bar_h = Inches(0.62)
bar_x0 = Inches(0.7)
bar_w_total = Inches(12.0)
total_aa = 210

domains = [
    (1, 30, "N 端延伸段\n(Ser23/24 磷酸化)", RGBColor(0x8A, 0x9B, 0xB8), False),
    (31, 136, "TnC / TnT 结合区", RGBColor(0xB9, 0xC8, 0xDC), False),
    (137, 148, "抑制区\n137–148", RED, True),
    (149, 163, "开关区\n149–163", ORANGE, True),
    (164, 210, "C 端移动区 164–210\n(173–181 结合肌动蛋白)", BLUE, True),
]

for start, end, label, color, hot in domains:
    x = bar_x0 + Emu(int(bar_w_total * (start - 1) / total_aa))
    w = Emu(int(bar_w_total * (end - start + 1) / total_aa))
    box = add_box(s, x, bar_y, w, bar_h, fill=color, shape=MSO_SHAPE.RECTANGLE, line=WHITE, line_w=Pt(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, seg in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = seg
        set_font(r, 9.5 if hot else 10.5, WHITE, True)

# hotspot bracket under bar
hot_x = bar_x0 + Emu(int(bar_w_total * 136 / total_aa))
hot_w = Emu(int(bar_w_total * (210 - 136) / total_aa))
add_box(s, hot_x, bar_y + bar_h + Inches(0.1), hot_w, Inches(0.05), fill=RED, shape=MSO_SHAPE.RECTANGLE)
add_text(s, Inches(5.5), bar_y + bar_h + Inches(0.2), Inches(7.4), Inches(0.35),
         [[("突变热点区（141–209）：抑制区 · 开关区 · 肌动蛋白结合区 · C 端保守肽段", 11.5, RED, True)]],
         align=PP_ALIGN.RIGHT)

add_text(s, Inches(0.7), Inches(6.75), Inches(12.0), Inches(0.5), [
    [("关键机制：", 12.5, NAVY, True),
     ("这些结构域共同决定“抑制—解除抑制”的钙调控开关；突变通过改变肌丝钙敏感性与抑制功能导致不同表型的心肌病。",
      12.5, GRAY, False)],
])

# ============================================================
# Slide 2 — Three phenotypes
# ============================================================
s = prs.slides.add_slide(blank)
header(s, "TNNI3 突变的三种心肌病表型", "One Gene, Three Phenotypes: HCM · RCM · DCM", 2)

cards = [
    ("HCM 肥厚型心肌病", RED, LIGHT_RED, "最常见（约 85% 的报道突变）", [
        ("遗传：", "常染色体显性错义突变，位于 C 端功能区"),
        ("代表突变：", "R145G/Q、S166F、R170W/Q、K183del、G203S、K206Q"),
        ("临床特点：", "心尖肥厚多见；肥厚程度和流出道梗阻较轻"),
        ("注意：", "猝死（SCD）风险偏高，可伴限制型生理表现"),
    ]),
    ("RCM 限制型心肌病", ORANGE, LIGHT_ORANGE, "TNNI3 是遗传性 RCM 首要致病基因", [
        ("遗传：", "显性错义突变，可为新发（de novo），多早发、重症"),
        ("代表突变：", "L144Q、R145W、A171T、K178E、D190G/H、R192H"),
        ("临床特点：", "限制性充盈、严重舒张功能障碍，无明显肥厚"),
        ("注意：", "同一家系内 HCM 与 RCM 可并存（同一疾病谱）"),
    ]),
    ("DCM 扩张型心肌病", TEAL, LIGHT_TEAL, "相对少见，分两种遗传模式", [
        ("显性错义：", "K36Q（N 端区）、N185K → 钙敏感性降低"),
        ("隐性双等位：", "纯合 A2V；双等位截短变异（如 R98*、R136*）"),
        ("临床特点：", "双等位截短 → 婴幼儿期（<2 岁）重症 DCM"),
        ("注意：", "常需心脏移植或早期死亡；预后差"),
    ]),
]

card_w = Inches(4.05)
gap = Inches(0.19)
x0 = Inches(0.45)
for i, (title, color, bg, tag, items) in enumerate(cards):
    x = x0 + (card_w + gap) * i
    add_box(s, x, Inches(1.4), card_w, Inches(5.0), fill=bg)
    add_box(s, x, Inches(1.4), card_w, Inches(0.62), fill=color, radius=0.18)
    add_text(s, x, Inches(1.5), card_w, Inches(0.45),
             [[(title, 17, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.25), Inches(2.18), card_w - Inches(0.5), Inches(0.5),
             [[(tag, 12, color, True)]])
    body = []
    for label, text in items:
        body.append([("▪ ", 11.5, color, True), (label, 11.5, color, True), (text, 11.5, DARK, False)])
    add_text(s, x + Inches(0.25), Inches(2.62), card_w - Inches(0.5), Inches(3.7), body, space_after=9)

add_box(s, Inches(0.45), Inches(6.55), Inches(12.45), Inches(0.72), fill=LIGHT_BLUE)
add_text(s, Inches(0.75), Inches(6.66), Inches(11.9), Inches(0.55), [
    [("表型重叠：", 12.5, BLUE, True),
     ("同一突变（如开关区 A157V）在同一家系中可分别表现为 HCM、RCM 甚至 DCM —— 提示遗传修饰因素与环境共同影响最终表型。",
      12.5, DARK, False)],
], anchor=MSO_ANCHOR.MIDDLE)

# ============================================================
# Slide 3 — Mechanism (calcium sensitivity spectrum) + clinical implications
# ============================================================
s = prs.slides.add_slide(blank)
header(s, "机制统一解释与临床意义", "The Calcium-Sensitivity Spectrum and Clinical Implications", 3)

add_text(s, Inches(0.45), Inches(1.3), Inches(12.0), Inches(0.4),
         [[("“钙敏感性连续谱”：功能改变方向决定表型", 16, NAVY, True)]])

rows = [
    ("钙敏感性中度升高", "HCM", RED, "抑制功能部分受损 → 收缩过度激活、舒张受损"),
    ("钙敏感性大幅升高\n（ΔpCa₅₀ 0.36–0.47）", "RCM", ORANGE, "无钙状态下不能完全松弛 → 严重舒张功能障碍而无明显肥厚"),
    ("钙敏感性降低\n最大 ATP 酶活性下降", "显性 DCM", TEAL, "收缩激活不足 → 收缩功能障碍、心腔扩大"),
    ("蛋白截短 / 功能丧失\n（双等位）", "隐性早发 DCM", NAVY, "cTnI 蛋白量显著减少 → 婴幼儿重症 DCM"),
]

ty = Inches(1.8)
row_h = Inches(0.78)
for i, (func, pheno, color, mech) in enumerate(rows):
    y = ty + Emu(int((row_h + Inches(0.1)) * i))
    add_box(s, Inches(0.45), y, Inches(3.5), row_h, fill=LIGHT_BLUE)
    tb = add_text(s, Inches(0.6), y, Inches(3.25), row_h,
                  [[(seg, 11.5, DARK, True)] for seg in func.split("\n")],
                  anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    # arrow
    ar = add_box(s, Inches(4.05), y + Inches(0.24), Inches(0.55), Inches(0.3),
                 fill=GRAY, shape=MSO_SHAPE.RIGHT_ARROW)
    add_box(s, Inches(4.7), y, Inches(2.0), row_h, fill=color, radius=0.15)
    add_text(s, Inches(4.7), y, Inches(2.0), row_h,
             [[(pheno, 14, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    add_text(s, Inches(6.9), y, Inches(6.1), row_h,
             [[(mech, 11.5, DARK, False)]], anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# Example callout
add_box(s, Inches(0.45), Inches(5.45), Inches(12.45), Inches(0.62), fill=LIGHT_ORANGE)
add_text(s, Inches(0.75), Inches(5.53), Inches(11.9), Inches(0.48), [
    [("典型例证：", 12, ORANGE, True),
     ("同一位点 145 —— R145G 致 HCM，R145W 致 RCM；钙敏感性升高幅度不同，表型即不同。R192H 丧失与原肌球蛋白结合，破坏 C 端“钙脱敏”功能。",
      12, DARK, False)],
], anchor=MSO_ANCHOR.MIDDLE)

# Clinical implications
add_box(s, Inches(0.45), Inches(6.15), Inches(12.45), Inches(1.2), fill=LIGHT_BLUE)
add_text(s, Inches(0.75), Inches(6.23), Inches(11.9), Inches(0.32),
         [[("临床意义", 13.5, BLUE, True)]])
add_text(s, Inches(0.75), Inches(6.57), Inches(11.9), Inches(0.75), [
    [("▪ ", 10.5, BLUE, True), ("特发性 RCM 患者应考虑 TNNI3 基因检测；家系筛查需同时关注 RCM 与 HCM 两种表型", 10.5, DARK, False)],
    [("▪ ", 10.5, BLUE, True), ("TNNI3-HCM 猝死风险偏高，应加强危险分层与随访", 10.5, DARK, False)],
    [("▪ ", 10.5, BLUE, True), ("婴幼儿重症 DCM 应注意双等位截短变异（隐性遗传，父母为无症状携带者）", 10.5, DARK, False)],
], space_after=3)

prs.save("/workspace/TNNI3基因突变与心肌病.pptx")
print("saved")
