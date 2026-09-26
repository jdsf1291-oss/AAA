#!/usr/bin/env python3
"""Generate 2 slides: thick vs thin filament HCM phenotype differences.
Based on Keyt LK et al. Front Cardiovasc Med. 2022;9:972301."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
BLUE = RGBColor(0x2E, 0x5E, 0xAA)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xFA)
RED = RGBColor(0xC0, 0x3A, 0x2B)
ORANGE = RGBColor(0xD9, 0x7B, 0x29)
TEAL = RGBColor(0x1F, 0x8A, 0x70)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_RED = RGBColor(0xFB, 0xEA, 0xE8)
LIGHT_ORANGE = RGBColor(0xFC, 0xF1, 0xE3)
LIGHT_TEAL = RGBColor(0xE4, 0xF4, 0xF0)
ROW_ALT = RGBColor(0xF4, 0xF7, 0xFC)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
blank = prs.slide_layouts[6]


def set_font(run, size, color=DARK, bold=False):
    f = run.font
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_box(slide, x, y, w, h, fill=None, line=None, line_w=None,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=3):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for line in lines:
        para_opts = {}
        if isinstance(line, tuple) and len(line) == 2 and isinstance(line[1], dict):
            runs, para_opts = line
        else:
            runs = line
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para_opts.get("align", align)
        p.space_after = Pt(para_opts.get("space_after", space_after))
        for text, size, color, bold in runs:
            r = p.add_run()
            r.text = text
            set_font(r, size, color, bold)
    return tb


def header(slide, title_cn, subtitle):
    add_box(slide, 0, 0, SW, Inches(0.92), fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    add_box(slide, 0, Inches(0.92), SW, Inches(0.05), fill=RED, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, Inches(0.5), Inches(0.1), Inches(12.3), Inches(0.5),
             [[(title_cn, 23, WHITE, True)]])
    add_text(slide, Inches(0.5), Inches(0.56), Inches(12.3), Inches(0.3),
             [[(subtitle, 11, RGBColor(0xB8, 0xC6, 0xDD), False)]])


def cite(slide):
    add_text(slide, Inches(0.35), Inches(7.14), Inches(12.6), Inches(0.3),
             [[("参考文献：Keyt LK, et al. Thin filament cardiomyopathies: A review of genetics, disease mechanisms, "
                "and emerging therapeutics. Front Cardiovasc Med. 2022;9:972301. doi:10.3389/fcvm.2022.972301",
                8.5, GRAY, False)]])


# ============================================================
# Slide 1 — clinical phenotype comparison
# ============================================================
s = prs.slides.add_slide(blank)
header(s, "粗肌丝 vs 细肌丝突变：HCM 临床表型差异",
       "Thick vs Thin Filament HCM: Distinct Clinical Phenotypes (Front Cardiovasc Med 2022;9:972301)")

# gene strip
strip_y = Inches(1.08)
add_box(s, Inches(0.35), strip_y, Inches(6.22), Inches(0.72), fill=LIGHT_BLUE)
add_text(s, Inches(0.6), strip_y + Inches(0.07), Inches(5.8), Inches(0.6), [
    [("粗肌丝基因：", 10.5, BLUE, True), ("MYH7（β-肌球蛋白重链）、MYBPC3（肌球蛋白结合蛋白 C）、MYL2/MYL3", 10.5, DARK, False)],
    [("—— 肌节突变中最常见、研究最充分", 9.5, GRAY, False)],
], anchor=MSO_ANCHOR.MIDDLE)
add_box(s, Inches(6.76), strip_y, Inches(6.22), Inches(0.72), fill=LIGHT_RED)
add_text(s, Inches(7.01), strip_y + Inches(0.07), Inches(5.8), Inches(0.6), [
    [("细肌丝基因：", 10.5, RED, True), ("TNNT2、TNNI3、TPM1、ACTC1、TNNC1", 10.5, DARK, False)],
    [("—— 细肌丝 HCM 以 TNNI3 和 TNNT2 最常见", 9.5, GRAY, False)],
], anchor=MSO_ANCHOR.MIDDLE)

# comparison table
tbl_x, tbl_y = Inches(0.35), Inches(1.98)
col_w = [Inches(2.55), Inches(4.95), Inches(5.13)]
header_h = Inches(0.4)
headers = ["特征", "粗肌丝 HCM", "细肌丝 HCM"]
head_colors = [NAVY, BLUE, RED]

rows = [
    ("肥厚程度与分布",
     "经典表现：非对称室间隔肥厚，程度较重",
     "肥厚程度较轻、分布不典型（心尖肥厚多见），起病年龄更轻"),
    ("流出道（LVOT）梗阻",
     "动态梗阻常见（二尖瓣前向运动明显）",
     "梗阻显著少见；室间隔减容治疗率更低"),
    ("舒张功能障碍",
     "存在，相对较轻",
     "更突出；更易出现限制型生理表现"),
    ("心衰进展",
     "进展相对缓慢",
     "进展至 NYHA III–IV 级的风险高 2 倍以上（独立于梗阻）；收缩功能障碍更重、出现更早（Coppini 等）"),
    ("心律失常",
     "房颤相对少",
     "房颤负荷高（导管消融率更高）；TNNT2 家系猝死率高（青年男性死亡率可达 64%）"),
    ("猝死（SCD）风险",
     "与肥厚程度相关性较大",
     "SCD 风险与肥厚程度不平行 —— 肥厚轻不等于风险低；不同研究结论有差异，依具体基因型而定"),
]

cx = tbl_x
for i, h in enumerate(headers):
    add_box(s, cx, tbl_y, col_w[i], header_h, fill=head_colors[i], shape=MSO_SHAPE.RECTANGLE, line=WHITE, line_w=Pt(1))
    add_text(s, cx + Inches(0.12), tbl_y, col_w[i] - Inches(0.24), header_h,
             [[(h, 11.5, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    cx += col_w[i]

row_h = Inches(0.64)
cy = tbl_y + header_h
for r, (feat, thick, thin) in enumerate(rows):
    bg = ROW_ALT if r % 2 == 0 else WHITE
    cx = tbl_x
    for c, txt in enumerate((feat, thick, thin)):
        add_box(s, cx, cy, col_w[c], row_h, fill=bg, shape=MSO_SHAPE.RECTANGLE,
                line=RGBColor(0xD5, 0xDE, 0xEA), line_w=Pt(0.75))
        color = NAVY if c == 0 else DARK
        add_text(s, cx + Inches(0.12), cy + Inches(0.03), col_w[c] - Inches(0.24), row_h - Inches(0.06),
                 [[(txt, 9.5, color, c == 0)]], anchor=MSO_ANCHOR.MIDDLE, space_after=1)
        cx += col_w[c]
    cy += row_h

# bottom link to case
add_box(s, Inches(0.35), cy + Inches(0.1), Inches(12.63), Inches(0.52), fill=LIGHT_ORANGE)
add_text(s, Inches(0.6), cy + Inches(0.16), Inches(12.2), Inches(0.42), [
    [("联系本例：", 10.5, ORANGE, True),
     ("TNNI3 属细肌丝基因 —— 本例肥厚不显著、无梗阻、以舒张/收缩功能障碍与心衰快速进展为主的“非经典”表现，符合细肌丝 HCM 的表型规律。",
      10.5, DARK, False)],
], anchor=MSO_ANCHOR.MIDDLE)

cite(s)

# ============================================================
# Slide 2 — mechanism differences & therapeutic implications
# ============================================================
s = prs.slides.add_slide(blank)
header(s, "粗肌丝 vs 细肌丝：致病机制差异与治疗启示",
       "Distinct Pathomechanisms: Myosin States vs Calcium Dysregulation, and Therapeutic Implications")

# two mechanism panels
py = Inches(1.12)
ph = Inches(3.35)
add_box(s, Inches(0.35), py, Inches(6.22), ph, fill=LIGHT_BLUE)
add_text(s, Inches(0.6), py + Inches(0.12), Inches(5.8), Inches(0.35),
         [[("粗肌丝 HCM：肌球蛋白“超收缩”机制", 14, BLUE, True)]])
add_text(s, Inches(0.6), py + Inches(0.55), Inches(5.75), ph - Inches(0.7), [
    [("▪ ", 10.5, BLUE, True), ("突变促进肌球蛋白头部与肌动蛋白相互作用，延长附着状态时间", 10.5, DARK, False)],
    [("▪ ", 10.5, BLUE, True), ("肌球蛋白由节能的超松弛态（SRX）转向无序松弛态（DRX）增多", 10.5, DARK, False)],
    [("▪ ", 10.5, BLUE, True), ("ATP 酶活性升高 → 高收缩性、耗氧增加 → 肥厚", 10.5, DARK, False)],
    [("▪ ", 10.5, BLUE, True), ("净效应：收缩力增强 + 舒张受损，以“动力过强”为核心", 10.5, DARK, False)],
], space_after=8)

add_box(s, Inches(6.76), py, Inches(6.22), ph, fill=LIGHT_RED)
add_text(s, Inches(7.01), py + Inches(0.12), Inches(5.8), Inches(0.35),
         [[("细肌丝 HCM：肌节钙失调机制", 14, RED, True)]])
add_text(s, Inches(7.01), py + Inches(0.55), Inches(5.75), ph - Inches(0.7), [
    [("▪ ", 10.5, RED, True), ("突变多位于关键调控域或蛋白相互作用界面（如 cTnI C 端）", 10.5, DARK, False)],
    [("▪ ", 10.5, RED, True), ("核心机制是肌节内钙调控失衡：钙敏感性升高、钙解离受损", 10.5, DARK, False)],
    [("▪ ", 10.5, RED, True), ("与粗肌丝相反：ATP 酶活性反而降低、肌球蛋白偏向 SRX，但氧耗与能量消耗仍增加 → 能量失衡", 10.5, DARK, False)],
    [("▪ ", 10.5, RED, True), ("净效应：舒张障碍突出、限制表型倾向、致心律失常性增高，肥厚反而较轻", 10.5, DARK, False)],
], space_after=8)

# therapy implications
ty = py + ph + Inches(0.15)
th = Inches(1.9)
add_box(s, Inches(0.35), ty, Inches(12.63), th, fill=LIGHT_TEAL)
add_text(s, Inches(0.6), ty + Inches(0.1), Inches(12.2), Inches(0.35),
         [[("治疗启示：机制不同 → 疗效可能不同", 14, TEAL, True)]])
add_text(s, Inches(0.6), ty + Inches(0.5), Inches(12.2), th - Inches(0.6), [
    [("▪ ", 10, TEAL, True),
     ("现有心衰/HCM 治疗（负性肌力药、室间隔减容）均未直接针对细肌丝的肌节功能异常；现行指南（含 ICD 指征）尚未区分粗/细肌丝基因型", 10, DARK, False)],
    [("▪ ", 10, TEAL, True),
     ("肌球蛋白 ATP 酶抑制剂 mavacamten（EXPLORER-HCM Ⅲ期试验证实改善梗阻性 HCM）理论上更契合粗肌丝机制；体外研究显示其也能逆转细肌丝突变（R92Q-cTnT、R145G-cTnI）的部分效应并降低胞浆钙水平，但对无肥厚/限制表型患者的获益未知", 10, DARK, False)],
    [("▪ ", 10, TEAL, True),
     ("提示未来需要基因型导向的个体化治疗；对细肌丝（TNNI3）患者，风险评估不应依赖肥厚程度，需综合心衰进展与心律失常负荷", 10, DARK, False)],
], space_after=5)

cite(s)

prs.save("/workspace/粗细肌丝HCM表型差异.pptx")
print("saved")
